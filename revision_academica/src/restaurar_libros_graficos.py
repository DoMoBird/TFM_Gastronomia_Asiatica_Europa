"""Preserve original literal-data workbooks lost by PPTX import/export, updating edited chart cells."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
from copy import deepcopy
import sys,re,posixpath
from lxml import etree as E
from openpyxl import load_workbook
source,candidate=map(Path,sys.argv[1:])
ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
with ZipFile(source) as z:old={n:z.read(n) for n in z.namelist()}
with ZipFile(candidate) as z:new={n:z.read(n) for n in z.namelist()}
for i in range(1,4):
 name=f'ppt/slides/charts/chart{i}.xml';relname=f'ppt/slides/charts/_rels/chart{i}.xml.rels'
 tree=E.fromstring(new[name]);orig=E.fromstring(old[name]);tree.append(deepcopy(orig.find('c:externalData',ns)))
 rel=E.fromstring(old[relname]);target=rel[0].get('Target');bookpath=posixpath.normpath(posixpath.join(posixpath.dirname(name),target))
 wb=load_workbook(BytesIO(old[bookpath]));ws=wb['Chart Data']
 assert not any(c.data_type=='f' for row in ws for c in row),'Original formulas must be preserved separately'
 cats=tree.xpath('//c:cat//c:pt/c:v/text()',namespaces=ns);values=tree.xpath('//c:val//c:pt/c:v/text()',namespaces=ns)
 assert len(cats)==len(values)
 for row,(cat,val) in enumerate(zip(cats,values),2):ws.cell(row,1,cat);ws.cell(row,2,float(val))
 b=BytesIO();wb.save(b);new[bookpath]=b.getvalue();new[relname]=old[relname];new[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8')
ct=E.fromstring(new['[Content_Types].xml']);oct=E.fromstring(old['[Content_Types].xml'])
for child in oct:
 if child.get('Extension')=='xlsx' or '/embeddings/' in child.get('PartName',''):
  if not any(x.attrib==child.attrib for x in ct):ct.append(deepcopy(child))
new['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8')
with ZipFile(candidate,'w',ZIP_DEFLATED) as z:
 for n,data in new.items():z.writestr(n,data)
